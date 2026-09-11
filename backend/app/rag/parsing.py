"""Document parsing: PDF / TXT / Markdown / DOCX -> cleaned text per page."""
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path

from app.core.errors import AppError

SUPPORTED_EXTENSIONS = {".pdf": "application/pdf", ".txt": "text/plain", ".md": "text/markdown",
                        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}


class ParsingError(AppError):
    status_code = 422
    code = "parsing_failed"


@dataclass
class ParsedPage:
    number: int
    text: str


@dataclass
class ParsedDocument:
    pages: list[ParsedPage]

    @property
    def page_count(self) -> int:
        return len(self.pages)


_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_DIGITS = re.compile(r"\d+")


def clean_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("­", "")
    text = _CONTROL.sub("", text)
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # re-join words hyphenated across line breaks
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def strip_running_lines(pages: list[ParsedPage]) -> list[ParsedPage]:
    """Remove running headers/footers: lines that (ignoring digits) repeat on most pages."""
    if len(pages) < 2:
        return pages

    def key(line: str) -> str:
        return _DIGITS.sub("#", line.strip())

    counts: dict[str, int] = {}
    for page in pages:
        for k in {key(ln) for ln in page.text.splitlines() if ln.strip()}:
            counts[k] = counts.get(k, 0) + 1
    running = {k for k, n in counts.items() if n >= max(2, int(len(pages) * 0.6))}
    out = []
    for p in pages:
        kept = [ln for ln in p.text.splitlines() if key(ln) not in running]
        # A "running line" that is the page's only content is body text, not a header/footer.
        out.append(ParsedPage(p.number, "\n".join(kept) if any(ln.strip() for ln in kept) else p.text))
    return out


def _parse_pdf(path: Path) -> list[ParsedPage]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise ParsingError("Encrypted PDFs are not supported")
        pages = [ParsedPage(i + 1, clean_text(page.extract_text() or "")) for i, page in enumerate(reader.pages)]
        return strip_running_lines(pages)
    except PdfReadError as exc:
        raise ParsingError(f"Could not read PDF: {exc}") from exc


def _parse_text(path: Path) -> list[ParsedPage]:
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    # Form feeds mark page breaks in plain-text exports.
    return [ParsedPage(i + 1, clean_text(part)) for i, part in enumerate(text.split("\f"))]


def _parse_docx(path: Path) -> list[ParsedPage]:
    import docx

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ParsingError(f"Could not read DOCX: {exc}") from exc
    lines = []
    for para in document.paragraphs:
        style = (para.style.name or "").lower() if para.style is not None else ""
        text = para.text.strip()
        if not text:
            continue
        if style == "title":
            lines += [f"# {text}", ""]
        elif style.startswith("heading"):
            level = int(style.split()[-1]) if style.split()[-1].isdigit() else 1
            lines += [f"{'#' * (level + 1)} {text}", ""]
        else:
            lines += [text, ""]
    return [ParsedPage(1, clean_text("\n".join(lines)))]  # DOCX has no fixed pagination


def parse_file(path: Path) -> ParsedDocument:
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ParsingError(f"Unsupported file type '{ext}'. Supported: {', '.join(SUPPORTED_EXTENSIONS)}")
    pages = _parse_pdf(path) if ext == ".pdf" else _parse_docx(path) if ext == ".docx" else _parse_text(path)
    if not any(p.text.strip() for p in pages):
        raise ParsingError("No extractable text found (scanned documents require OCR, which is not supported)")
    return ParsedDocument(pages)
