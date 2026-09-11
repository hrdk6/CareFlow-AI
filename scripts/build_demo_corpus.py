"""Render the synthetic knowledge base (rag/corpus/source/*.md) into the upload formats listed in
rag/corpus/manifest.json: PDF (fpdf2), DOCX (python-docx) and TXT. Output: rag/corpus/dist/.

    uv run --project backend python scripts/build_demo_corpus.py
"""
import json
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
CORPUS = ROOT / "rag" / "corpus"
FOOTER = "SYNTHETIC DEMO DOCUMENT - CareFlow AI portfolio project - not clinical guidance"


def blocks(md: str) -> list[tuple[str, str]]:
    """[(kind, text)] where kind in title/h1/h2/p."""
    out = []
    for para in re.split(r"\n\s*\n", md.strip()):
        para = " ".join(line.strip() for line in para.splitlines())
        if para.startswith("### "):
            out.append(("h2", para[4:]))
        elif para.startswith("## "):
            out.append(("h1", para[3:]))
        elif para.startswith("# "):
            out.append(("title", para[2:]))
        else:
            out.append(("p", para))
    return out


def latin1(text: str) -> str:
    return text.replace("’", "'").replace("–", "-").replace("—", "-").encode("latin-1", "replace").decode("latin-1")


def to_pdf(md: str, target: Path) -> None:
    from fpdf import FPDF

    class Doc(FPDF):
        def footer(self):
            self.set_y(-12)
            self.set_font("Helvetica", "I", 7)
            self.cell(0, 8, f"{FOOTER} - page {self.page_no()}", align="C")

    pdf = Doc(format="A4")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_margins(20, 18, 20)
    pdf.add_page()
    styles = {"title": ("B", 15, 9), "h1": ("B", 12.5, 7), "h2": ("B", 11, 6), "p": ("", 10.5, 5.4)}
    for kind, text in blocks(md):
        style, size, height = styles[kind]
        if kind in ("h1", "h2"):
            pdf.ln(2)
        pdf.set_font("Helvetica", style, size)
        pdf.multi_cell(0, height, latin1(text))
        pdf.ln(2 if kind == "p" else 1)
    pdf.output(str(target))


def to_docx(md: str, target: Path) -> None:
    import docx

    document = docx.Document()
    for kind, text in blocks(md):
        if kind == "title":
            document.add_heading(text, level=0)
        elif kind == "h1":
            document.add_heading(text, level=1)
        elif kind == "h2":
            document.add_heading(text, level=2)
        else:
            document.add_paragraph(text)
    document.save(str(target))


def to_txt(md: str, target: Path) -> None:
    lines = []
    for _kind, text in blocks(md):
        lines.append(text)
        lines.append("")
    target.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    manifest = json.loads((CORPUS / "manifest.json").read_text())
    dist = CORPUS / "dist"
    dist.mkdir(exist_ok=True)
    for entry in manifest["documents"]:
        src = CORPUS / "source" / entry["source"]
        target = dist / f"{src.stem}.{entry['format']}"
        md = src.read_text(encoding="utf-8")
        {"pdf": to_pdf, "docx": to_docx, "txt": to_txt}[entry["format"]](md, target)
        print(f"built {target.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
